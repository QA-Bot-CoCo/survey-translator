import io
import json
import unittest
import zipfile
from unittest.mock import Mock, patch

import openpyxl
from openpyxl.styles import Font

from survey.settings import Settings
from survey.translation import BRANDS, OpenAITranslator, TranslationError, system_prompt
from survey.web import create_app
from survey.workbooks import InvalidWorkbook, extract_cells, make_batches, render_workbook, validate_archive

ORIGIN = "https://survey.test"


def settings():
    return Settings(api_key="test-only", model="gpt-6-luna", origin=ORIGIN)


def workbook(content="Hello", messages="Welcome"):
    wb = openpyxl.Workbook()
    first = wb.active
    first.title = "Survey Content"
    first.append(["ID", "Content", "Translated Content"])
    first.append(["intro", content, None])
    first["B2"].font = Font(bold=True)
    second = wb.create_sheet("Survey Messages")
    second.append(["ID", "Messages", "Translated Messages"])
    second.append(["success", messages, None])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def amend_archive(data, name, value):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as old, zipfile.ZipFile(out, "w") as new:
        for info in old.infolist():
            if info.filename != name:
                new.writestr(info.filename, old.read(info.filename))
        new.writestr(name, value)
    return out.getvalue()


class WorkbookTests(unittest.TestCase):
    def test_both_sheets_are_translated_and_styles_preserved(self):
        data = workbook()
        cells, columns = extract_cells(data)
        self.assertEqual([cell["text"] for cell in cells], ["Hello", "Welcome"])
        result = render_workbook(data, cells, columns, {"item0": "Hallo", "item1": "Willkommen"})
        wb = openpyxl.load_workbook(io.BytesIO(result))
        self.assertEqual(wb["Survey Content"]["B2"].value, "Hello")
        self.assertEqual(wb["Survey Content"]["C2"].value, "Hallo")
        self.assertEqual(wb["Survey Messages"]["C2"].value, "Willkommen")
        self.assertTrue(wb["Survey Content"]["C2"].font.bold)
        wb.close()

    def test_model_formulas_are_literal_strings(self):
        data = workbook()
        cells, columns = extract_cells(data)
        result = render_workbook(data, cells, columns, {"item0": "=1+1", "item1": "@SUM(A1)"})
        wb = openpyxl.load_workbook(io.BytesIO(result))
        self.assertEqual(wb["Survey Content"]["C2"].data_type, "s")
        self.assertEqual(wb["Survey Content"]["C2"].value, "=1+1")
        wb.close()

    def test_rejects_source_formulas(self):
        with self.assertRaises(InvalidWorkbook):
            extract_cells(workbook("=1+1"))

    def test_rejects_zip_traversal(self):
        with self.assertRaises(InvalidWorkbook):
            validate_archive(amend_archive(workbook(), "../outside.xml", b"<x/>"))

    def test_rejects_macros(self):
        with self.assertRaises(InvalidWorkbook):
            validate_archive(amend_archive(workbook(), "xl/vbaProject.bin", b"macro"))

    def test_rejects_external_relationship(self):
        xml = b'<Relationships><Relationship TargetMode="External" Target="file:///secret"/></Relationships>'
        with self.assertRaises(InvalidWorkbook):
            validate_archive(amend_archive(workbook(), "xl/_rels/evil.rels", xml))

    def test_rejects_dtd_and_entities(self):
        xml = b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///secret">]><x>&secret;</x>'
        with self.assertRaises(InvalidWorkbook):
            validate_archive(amend_archive(workbook(), "custom.xml", xml))

    def test_rejects_expansion_bomb(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("xl/workbook.xml", "<x>" + "x" * 100000 + "</x>")
        with self.assertRaises(InvalidWorkbook):
            validate_archive(out.getvalue())

    def test_rejects_non_zip(self):
        with self.assertRaises(InvalidWorkbook):
            extract_cells(b"not a workbook")

    def test_rejects_missing_messages(self):
        wb = openpyxl.Workbook()
        wb.active.append(["Content"])
        wb.active.append(["Hello"])
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        with self.assertRaises(InvalidWorkbook):
            extract_cells(out.getvalue())

    def test_rejects_oversized_cell(self):
        with self.assertRaises(InvalidWorkbook):
            extract_cells(workbook("a" * 2001))

    def test_rejects_cell_outside_forged_dimensions(self):
        xml = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1"/><sheetData><row r="2"><c r="XFD2"><v>1</v></c></row></sheetData></worksheet>'
        with self.assertRaises(InvalidWorkbook):
            extract_cells(amend_archive(workbook(), "xl/worksheets/sheet1.xml", xml))

    def test_rejects_xlm_macro_sheet(self):
        with self.assertRaises(InvalidWorkbook):
            validate_archive(amend_archive(workbook(), "xl/macrosheets/sheet1.xml", b"<sheet/>"))

    def test_rejects_incomplete_translation(self):
        data = workbook()
        cells, columns = extract_cells(data)
        with self.assertRaises(InvalidWorkbook):
            render_workbook(data, cells, columns, {"item0": "Hallo"})

    def test_batches_bound_text_and_cell_counts(self):
        cells = [{"id": f"item{i}", "text": "x" * 1500} for i in range(17)]
        batches = make_batches(cells)
        self.assertEqual(sum(len(batch) for batch in batches), 17)
        self.assertTrue(all(sum(len(cell["text"]) for cell in batch) <= 6000 for batch in batches))


class ProviderTests(unittest.TestCase):
    def test_provider_uses_fixed_endpoint_and_strict_schema(self):
        response = Mock(status_code=200)
        response.json.return_value = {
            "choices": [{"finish_reason": "stop", "message": {"content": '{"item0":"Hallo"}'}}],
        }
        with patch("survey.translation.requests.post", return_value=response) as post:
            result = OpenAITranslator(settings()).translate(
                [{"id": "item0", "text": "Hello"}], "gaminator", "deDE",
            )
        self.assertEqual(result, {"item0": "Hallo"})
        self.assertEqual(post.call_args.args[0], "https://api.openai.com/v1/chat/completions")
        payload = post.call_args.kwargs["json"]
        self.assertFalse(payload["store"])
        self.assertTrue(payload["response_format"]["json_schema"]["strict"])
        self.assertFalse(post.call_args.kwargs["allow_redirects"])

    def test_incomplete_model_output_is_rejected(self):
        response = Mock(status_code=200)
        response.json.return_value = {
            "choices": [{"finish_reason": "stop", "message": {"content": "{}"}}],
        }
        with patch("survey.translation.requests.post", return_value=response):
            with self.assertRaises(TranslationError):
                OpenAITranslator(settings()).translate(
                    [{"id": "item0", "text": "Hello"}], "gaminator", "deDE",
                )

    def test_upstream_errors_do_not_disclose_response_body(self):
        response = Mock(status_code=401, text="secret-provider-response")
        with patch("survey.translation.requests.post", return_value=response):
            with self.assertRaises(TranslationError) as error:
                OpenAITranslator(settings()).translate(
                    [{"id": "item0", "text": "Hello"}], "gaminator", "deDE",
                )
        self.assertNotIn("secret-provider-response", str(error.exception))
